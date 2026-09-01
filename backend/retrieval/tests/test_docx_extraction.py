"""
Tests for .docx text extraction and for the "stored but unsearchable" reporting.

WHY. `_extract_docx` read only `document.paragraphs`, so a document whose text
lives in a TABLE extracted to zero characters. `upload_document` then took the
`if not chunks` branch and returned status "stored" with no log line, which is
the same thing it returns for an image, so nothing anywhere distinguished "this
is a picture" from "the parser could not read your file". Measured on the live
store: 6 of one user's 7 .docx uploads, every one a valid OOXML file with no
paragraph text and a single table holding 187 to 3636 characters, sat
unsearchable for three months and were reported to their owner as successful
uploads. Corpus-wide those 6 were the ONLY recoverable losses, so the blast
radius was small; the silence is what made it survive.

Two properties are pinned here, and the second matters more than the first:
extraction has to read tables, and an extraction that yields nothing has to say
so. A parser can always meet a shape it does not handle (text boxes, headers, a
scanned PDF); the requirement is that the next one announces itself on the first
upload instead of on a complaint months later.

    docker exec munin-retrieval python /app/tests/test_docx_extraction.py
"""

from __future__ import annotations

import asyncio
import importlib
import io
import sys
import traceback

sys.path.insert(0, "/app")
sys.path.insert(0, ".")

DS = importlib.import_module("document_store")


def _check(name, ok, detail=""):
    print(("[PASS] " if ok else "[FAIL] ") + name + (f"  {detail}" if detail and not ok else ""))
    return ok


def _docx_bytes(build) -> bytes:
    """Render a python-docx Document built by `build` to bytes."""
    import docx
    document = docx.Document()
    build(document)
    buf = io.BytesIO()
    document.save(buf)
    return buf.getvalue()


# --- extraction -------------------------------------------------------------

def test_table_only_document_is_extracted():
    """The exact shape that was lost: no paragraph text, one table."""
    def build(d):
        t = d.add_table(rows=2, cols=2)
        t.cell(0, 0).text = "Sample ID"
        t.cell(0, 1).text = "Measurement"
        t.cell(1, 0).text = "PM6:Y6"
        t.cell(1, 1).text = "18.4% PCE"
    text = DS._extract_docx(_docx_bytes(build))
    return _check("a table-only .docx yields its cell text",
                  "PM6:Y6" in text and "18.4% PCE" in text and "Sample ID" in text,
                  f"got {text!r}")


def test_paragraph_document_still_works():
    """The one file that DID extract must not regress."""
    def build(d):
        d.add_paragraph("First paragraph.")
        d.add_paragraph("Second paragraph.")
    text = DS._extract_docx(_docx_bytes(build))
    return _check("a paragraph .docx still extracts",
                  "First paragraph." in text and "Second paragraph." in text,
                  f"got {text!r}")


def test_document_order_is_preserved():
    """Concatenating all paragraphs then all tables would put the table after
    the closing prose, so chunking would glue a table to unrelated text."""
    def build(d):
        d.add_paragraph("INTRO")
        t = d.add_table(rows=1, cols=1)
        t.cell(0, 0).text = "TABLE"
        d.add_paragraph("OUTRO")
    text = DS._extract_docx(_docx_bytes(build))
    order = [text.find("INTRO"), text.find("TABLE"), text.find("OUTRO")]
    return _check("paragraphs and tables come out in document order",
                  -1 not in order and order == sorted(order), f"positions {order}")


def test_merged_cells_are_not_repeated():
    """`row.cells` yields the same cell once per grid column it spans, so a
    merged heading would otherwise be duplicated across the row."""
    def build(d):
        t = d.add_table(rows=1, cols=3)
        merged = t.cell(0, 0).merge(t.cell(0, 2))
        merged.text = "SPAN"
    text = DS._extract_docx(_docx_bytes(build))
    return _check("a horizontally merged cell appears once",
                  text.count("SPAN") == 1, f"count={text.count('SPAN')} in {text!r}")


def test_empty_cells_do_not_produce_noise():
    def build(d):
        t = d.add_table(rows=2, cols=2)
        t.cell(0, 0).text = "only"
    text = DS._extract_docx(_docx_bytes(build))
    return _check("blank cells and blank rows are dropped",
                  text.strip() == "only", f"got {text!r}")


def test_garbage_bytes_degrade_quietly():
    """A corrupt upload must return empty, not raise into the request."""
    return _check("non-docx bytes return empty rather than raising",
                  DS._extract_docx(b"this is not a zip file") == "")


def test_extracted_table_text_survives_chunking():
    """Extraction is only half the path: the text has to reach chunk_text as
    something non-empty, or the upload still lands as 'stored'."""
    def build(d):
        t = d.add_table(rows=1, cols=2)
        t.cell(0, 0).text = "Parameter"
        t.cell(0, 1).text = "Value measured at 300 K under one sun."
    chunks = DS.chunk_text(DS._extract_docx(_docx_bytes(build)))
    return _check("table text produces at least one chunk", len(chunks) >= 1,
                  f"{len(chunks)} chunks")


# --- the reporting half -----------------------------------------------------

class _FakeQdrant:
    def __init__(self):
        self.points = []

    def upsert(self, collection_name, points):
        self.points.extend(points)


class _FakeBGE:
    def encode(self, chunks, show_progress_bar=False):
        # upload_document calls .tolist() on the result, so return the ndarray
        # the real encoder returns rather than a plain list.
        import numpy as np
        return np.zeros((len(chunks), 768), dtype="float32")


def _run_upload(tmpdir, filename, payload, qdrant=None, bge=None):
    """Call upload_document with the index stubbed and the store redirected."""
    orig = (DS.USER_DOCS_DIR, DS.get_qdrant, DS.get_bge, DS.ensure_collection)
    DS.USER_DOCS_DIR = tmpdir
    DS.get_qdrant = lambda: qdrant
    DS.get_bge = lambda: bge
    DS.ensure_collection = lambda: True
    try:
        return asyncio.run(DS.upload_document(
            filename=filename, file_bytes=payload, user_email="u@example.org"))
    finally:
        (DS.USER_DOCS_DIR, DS.get_qdrant, DS.get_bge, DS.ensure_collection) = orig


def test_unreadable_text_upload_reports_a_reason():
    """"stored" alone is what let this hide: it is also what an image returns."""
    import tempfile
    def build(d):
        d.add_paragraph("")
    with tempfile.TemporaryDirectory() as tmp:
        res = _run_upload(tmp, "empty.docx", _docx_bytes(build),
                          qdrant=_FakeQdrant(), bge=_FakeBGE())
    return _check("an unreadable text upload says why it was not embedded",
                  res["status"] == "stored" and res.get("reason") == "no_text_extracted",
                  f"{res}")


def test_index_unavailable_is_a_distinct_reason():
    """A readable document plus a missing index is an OPERATIONAL failure, and
    conflating it with a parser failure sends the next person to the wrong
    place."""
    import tempfile
    def build(d):
        d.add_paragraph("Readable content that chunks fine.")
    with tempfile.TemporaryDirectory() as tmp:
        res = _run_upload(tmp, "fine.docx", _docx_bytes(build),
                          qdrant=None, bge=None)
    return _check("a missing index is reported as index_unavailable",
                  res["status"] == "stored" and res.get("reason") == "index_unavailable",
                  f"{res}")


def test_successful_upload_carries_no_reason():
    import tempfile
    def build(d):
        t = d.add_table(rows=1, cols=1)
        t.cell(0, 0).text = "Table content that should now embed."
    with tempfile.TemporaryDirectory() as tmp:
        res = _run_upload(tmp, "table.docx", _docx_bytes(build),
                          qdrant=_FakeQdrant(), bge=_FakeBGE())
    return _check("a table-only .docx now uploads as embedded",
                  res["status"] == "embedded" and res["chunks"] >= 1
                  and "reason" not in res, f"{res}")


def test_image_upload_is_unchanged():
    """Images legitimately store without embedding, and must not acquire a
    reason that reads like a defect."""
    import tempfile
    with tempfile.TemporaryDirectory() as tmp:
        res = _run_upload(tmp, "shot.png", b"\x89PNG\r\n\x1a\n" + b"0" * 64,
                          qdrant=_FakeQdrant(), bge=_FakeBGE())
    return _check("an image still stores cleanly with no reason",
                  res["status"] == "stored" and "reason" not in res, f"{res}")


TESTS = [
    test_table_only_document_is_extracted,
    test_paragraph_document_still_works,
    test_document_order_is_preserved,
    test_merged_cells_are_not_repeated,
    test_empty_cells_do_not_produce_noise,
    test_garbage_bytes_degrade_quietly,
    test_extracted_table_text_survives_chunking,
    test_unreadable_text_upload_reports_a_reason,
    test_index_unavailable_is_a_distinct_reason,
    test_successful_upload_carries_no_reason,
    test_image_upload_is_unchanged,
]


def main() -> int:
    failed = 0
    for fn in TESTS:
        try:
            if not fn():
                failed += 1
        except Exception:
            failed += 1
            print(f"[FAIL] {fn.__name__} raised:")
            traceback.print_exc()
    print(f"\n{len(TESTS) - failed}/{len(TESTS)} passed")
    return 1 if failed else 0


if __name__ == "__main__":
    sys.exit(main())
